import copy, re, sys, io
from pptx import Presentation
S=sys.argv[1]
dst=Presentation(S+'/v9/out.pptx'); src=Presentation(S+'/build/grupo10.pptx')
orig=list(dst.slides)
lay=[l for l in dst.slide_layouts if l.name=='DEFAULT'][0]
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
P='{http://schemas.openxmlformats.org/presentationml/2006/main}'
new=[]
for ss in src.slides:
    d=dst.slides.add_slide(lay)
    for ph in list(d.placeholders): ph._element.getparent().remove(ph._element)
    for el in ss.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
        e=copy.deepcopy(el)
        for b in e.iter():
            rid=b.get(R+'embed')
            if rid:
                _,nr=d.part.get_or_add_image_part(io.BytesIO(ss.part.related_part(rid).blob)); b.set(R+'embed',nr)
        d.shapes._spTree.append(e)
    bg=ss._element.cSld.find(P+'bg')
    if bg is not None: d._element.cSld.insert(0,copy.deepcopy(bg))
    tf=d.notes_slide.notes_text_frame
    if tf is not None: tf.text=ss.notes_slide.notes_text_frame.text
    new.append(d)
def title(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() and not sh.text_frame.text.startswith('CAPÍTULO'): return sh.text_frame.text
    return ''
k=[i for i,s in enumerate(orig) if title(s).startswith('Grupo de Desenvolvimento de Confinadores')]
assert len(k)==1, k
k=k[0]
order=orig[:k]+new+orig[k+1:]
lst=dst.slides._sldIdLst
el_by={el.get('id'):el for el in lst}
oid=[str(s.slide_id) for s in order]; rem=el_by[str(orig[k].slide_id)]
dst.part.drop_rel(rem.get(R+'id')); lst.remove(rem)
for el in list(lst): lst.remove(el)
for x in oid: lst.append(el_by[x])
H=914400
def set_para(p,text):
    runs=p.runs; runs[0].text=text
    for x in runs[1:]: x._r.getparent().remove(x._r)
for i,s in enumerate(order,1):
    for sh in s.shapes:
        if sh.has_text_frame and re.fullmatch(r'\d{1,2}',sh.text_frame.text.strip()) and sh.left is not None and sh.left>8.5*H and sh.top>4.9*H and sh.text_frame.paragraphs[0].runs:
            set_para(sh.text_frame.paragraphs[0],str(i))
dst.save(S+'/v10/out.pptx'); print('ok',len(order),'grupo em',k+1)
