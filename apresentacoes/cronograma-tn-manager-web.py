"""Gera o Cronograma TN Manager Web reestruturado (versao 09/10/2026)."""
import sys
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "Cronograma_TN_Manager_Web_v2.xlsx"
D = lambda d, m, y=2026: date(y, m, d)

NAVY, MAG, TEAL, LIME, GRAY, LIGHT = "1B3A8C", "C8106E", "0096B4", "C5D300", "D9DEE3", "F4F6F8"
PHASE_COLOR = {
    "Base Estratégica": "1B3A8C", "Imersão Técnica": "26668F", "Diagnóstico": "C8106E",
    "Produto e Direcionamento": "0096B4", "Consolidação da Estratégia": "00A8A8",
    "Validações": "6B4FA0", "Piloto": "7A858F",
}
thin = Side(style="thin", color=GRAY)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")

# item, item original, tópico, tarefa, objetivo, responsável, entregável, facilitador, início prev, início real, término prev, status
ROWS = [
    ("2.1", "2.1", "Base Estratégica", "Análise de Business Case TN Manager e TN Manager Web",
     "Clarificar onde o produto gera valor e para quem. Definir uma versão web mais simples, com lançamento manual (TN Manager Web), começando com poucas propriedades e evoluindo por versões.",
     "Rodrigo", "PPT do Business Case v2.0 (hipóteses): cliente-alvo, proposta de valor, produto mínimo e roadmap macro",
     "Equipe Técnica, Dreamer", D(1, 9), D(7, 9), D(16, 10), "Em andamento"),
    ("2.1.1", "2.1.1", "Base Estratégica", "Entrevista com stakeholders (interno) e benchmark com usuários de sistemas concorrentes (externo)",
     "Aderência interna: medir entendimento e apoio entre quem pode acelerar ou travar o projeto na Trouw. Benchmark externo: pontos fortes e fracos dos concorrentes, com auditorias que usam vários sistemas e consultores com sistema próprio.",
     "Rodrigo", "PPTX com análise do TN Manager versus principais concorrentes",
     "Equipe Técnica, fazendas, KJS Auditoria, consultores parceiros", D(14, 9), D(14, 9), D(23, 10), "Em andamento"),
    ("2.2.1", "Formulário (sem número)", "Base Estratégica", "Rascunho do formulário de pesquisa",
     "Montar a primeira versão do formulário para confinamentos clientes e não clientes, para a consultora revisar já com material pronto.",
     "Rodrigo", "Rascunho do formulário", "Equipe Técnica", D(19, 10), None, D(23, 10), "A iniciar"),
    ("2.2.2", "Formulário (sem número)", "Base Estratégica", "Revisão do formulário com a consultora Kátia Sato e teste com 2 ou 3 fazendas",
     "Revisar perguntas e formato com a consultora, testar com poucas fazendas e ajustar. Kátia retorna de férias na última semana de outubro; 02/11 é feriado (Finados).",
     "Rodrigo", "Formulário final validado", "Kátia Sato, Equipe Técnica", D(26, 10), None, D(6, 11), "A iniciar"),
    ("2.2.3", "2.2", "Base Estratégica", "Discovery com clientes (incluindo detratores): aplicação do formulário",
     "Enviar para clientes Trouw, ex-clientes e não clientes. Meta: mínimo de 30 fazendas respondendo (inteligência competitiva de mercado).",
     "Rodrigo", "Base de respostas (mínimo de 30 fazendas)", "Equipe Técnica, RTVs, fazendas", D(9, 11), None, D(27, 11), "A iniciar"),
    ("2.2.4", "2.2", "Base Estratégica", "Consolidação do discovery",
     "Entender a percepção real de valor e as dores antes de construir. Ranking das principais dores, com evidências.",
     "Rodrigo", "PPTX com conclusão das principais dores dos entrevistados", "Equipe Técnica", D(30, 11), None, D(4, 12), "A iniciar"),
    ("3.1", "3.1 + 3.1.1 + Visita Dreamer", "Imersão Técnica", "Imersão técnica (Software House), com visita à Dreamer e teste da entrega do TN Manager Web",
     "Entender arquitetura e limitações do TN Manager (Plano A). Analisar o documento de design do TN Manager Web (78 de 131 tarefas concluídas) e testar a entrega prevista para a última semana de novembro. Mapear alternativas (Plano B, ex.: Control Bell) com prazo, custo e aderência.",
     "Rodrigo", "Relatório técnico: estado atual, limitações e alternativas", "Matheus e Dreamer", D(16, 11), None, D(27, 11), "A iniciar"),
    ("4.1", "4.1 + 4.2 + 4.3", "Diagnóstico", "Diagnóstico consolidado e apresentação (SWOT, mercado e parte técnica)",
     "Juntar os resultados de 2.1.1, 2.2 e 3.1 em uma SWOT. Validar os pontos e definir o rumo do projeto. Checkpoint formal com a liderança.",
     "Rodrigo", "Apresentação do Diagnóstico com decisão de rumo", "Matheus, Equipe Técnica", D(7, 12), None, D(11, 12), "A iniciar"),
    ("5.1", "3.2", "Produto e Direcionamento", "Geração de valor e revisão de backlog (PRDs)",
     "Priorizar o que realmente gera valor, com base no Diagnóstico. Escrever os PRDs que vão direcionar a Dreamer nas próximas etapas.",
     "Rodrigo", "PRDs priorizados (Product Requirement Documents)", "Matheus e Dreamer", D(14, 12), None, D(18, 12), "A iniciar"),
    ("6.1", "5.1", "Consolidação da Estratégia", "Construção do Roadmap (2 anos)",
     "Definir a trajetória de evolução do TN Manager Web por versões, priorizando maior impacto em valor para o cliente e retorno para o negócio. Inclui plano de piloto com tamanho único e critérios de avanço.",
     "Rodrigo", "Roadmap de 24 meses", "Matheus e Dreamer", D(5, 1, 2027), None, D(15, 1, 2027), "A iniciar"),
    ("6.2", "5.2", "Consolidação da Estratégia", "Estruturação do GTM",
     "Escopo final definido com a liderança após a validação do Diagnóstico (4.1): público-alvo, proposta de valor, canais e abordagem comercial.",
     "Rodrigo", "Plano de GTM", "Liderança comercial", D(18, 1, 2027), None, D(29, 1, 2027), "A iniciar"),
    ("6.3", "5.3", "Consolidação da Estratégia", "Estratégia de Monetização",
     "Modelo final definido com a liderança, com base no que for validado no Diagnóstico (4.1).",
     "Rodrigo", "Modelo de monetização", "Liderança, Finanças", D(18, 1, 2027), None, D(29, 1, 2027), "A iniciar"),
    ("7.0", "6.0", "Validações", "Apresentação da BP e GTM",
     "Validar os entregáveis 6.1 a 6.3 e a decisão entre seguir com o desenvolvimento próprio ou migrar para parceria, com possibilidade de renegociar escopo, prazo ou investimento.",
     "Rodrigo", "Apresentação final (BP e GTM)", "Liderança", D(1, 2, 2027), None, D(5, 2, 2027), "A iniciar"),
    ("8.0", "Ideia amarela (Grupo de Desenvolvimento)", "Piloto", "Piloto com fazendas (Grupo de Desenvolvimento)",
     "Começar com poucas propriedades (sugestão: 5) e ampliar por etapas, só avançando com critérios cumpridos. Tamanho e data definidos no Roadmap (6.1), alinhados à ocupação dos confinamentos.",
     "Rodrigo", "Plano de piloto com metas e critérios de avanço", "Equipe Técnica, fazendas", None, None, None, "A definir no 6.1"),
]

HEAD = ["Item", "Item original", "Tópico", "Tarefas", "Objetivo", "Responsável", "Entregável", "Facilitador - TN",
        "Início Previsto", "Início Real", "Término Previsto", "Término Real", "Status"]
WIDTHS = [7, 16, 18, 34, 60, 12, 30, 24, 12, 12, 12, 12, 14]

wb = Workbook()
ws = wb.active
ws.title = "Cronograma-2026"
ws.sheet_view.showGridLines = False
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
ws["A1"] = "Cronograma: TN Manager Web"
ws["A1"].font = Font(name="Calibri", size=18, bold=True, color=NAVY)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 32
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEAD))
ws["A2"] = "Versão reestruturada em 09/10/2026. Coluna \"Item original\" mostra de onde veio cada linha da versão anterior."
ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="7A858F")
ws["A2"].alignment = Alignment(horizontal="center")

for c, (h, w) in enumerate(zip(HEAD, WIDTHS), 1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    cell.border = BORDER
    ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[4].height = 30

STATUS_FILL = {"Em andamento": "FFF2CC", "A iniciar": "FFFFFF", "A definir no 6.1": "EEEEEE"}
for r, row in enumerate(ROWS, 5):
    item, orig, top, tar, obj, resp, ent, fac, ini, inir, fim, st = row
    vals = [item, orig, top, tar, obj, resp, ent, fac, ini, inir, fim, None, st]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name="Calibri", size=10)
        cell.border = BORDER
        cell.alignment = CENTER if c in (1, 2, 6, 9, 10, 11, 12, 13) else WRAP
        if isinstance(v, date):
            cell.number_format = "DD/MM/YYYY"
    col = PHASE_COLOR[top]
    ws.cell(row=r, column=1).font = Font(name="Calibri", size=11, bold=True, color=col)
    ws.cell(row=r, column=3).font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor=col)
    ws.cell(row=r, column=4).font = Font(name="Calibri", size=10, bold=True)
    ws.cell(row=r, column=13).fill = PatternFill("solid", fgColor=STATUS_FILL[st])
    if r % 2 == 0:
        for c in (2, 4, 5, 6, 7, 8, 9, 10, 11, 12):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=LIGHT)
    ws.row_dimensions[r].height = 78
ws.freeze_panes = "D5"
ws.auto_filter.ref = f"A4:{get_column_letter(len(HEAD))}{4 + len(ROWS)}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

# Macro Cronograma (Gantt semanal)
g = wb.create_sheet("Macro Cronograma-2026")
g.sheet_view.showGridLines = False
start = D(31, 8)  # segunda-feira
weeks = [start + timedelta(weeks=i) for i in range(24)]
g.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3 + len(weeks))
g["A1"] = "Macro Cronograma: TN Manager Web (semanas começando na segunda-feira)"
g["A1"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
g.row_dimensions[1].height = 28
for c, h in enumerate(["Item", "Atividade", "Tópico"], 1):
    cell = g.cell(row=3, column=c, value=h)
    cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    cell.border = BORDER
g.column_dimensions["A"].width = 7
g.column_dimensions["B"].width = 46
g.column_dimensions["C"].width = 22
for i, wk in enumerate(weeks):
    c = 4 + i
    cell = g.cell(row=3, column=c, value=wk)
    cell.number_format = "DD/MM"
    cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
    cell.border = BORDER
    g.column_dimensions[get_column_letter(c)].width = 4.2
g.row_dimensions[3].height = 44
for r, row in enumerate(ROWS, 4):
    item, _, top, tar, *_rest = row
    ini, fim = row[8], row[10]
    g.cell(row=r, column=1, value=item).font = Font(name="Calibri", size=10, bold=True, color=PHASE_COLOR[top])
    g.cell(row=r, column=2, value=tar).font = Font(name="Calibri", size=10)
    g.cell(row=r, column=3, value=top).font = Font(name="Calibri", size=9, color="4A5560")
    for c in (1, 2, 3):
        g.cell(row=r, column=c).border = BORDER
        g.cell(row=r, column=c).alignment = WRAP if c == 2 else CENTER
    for i, wk in enumerate(weeks):
        cell = g.cell(row=r, column=4 + i)
        cell.border = BORDER
        if ini and fim and wk <= fim and wk + timedelta(days=6) >= ini:
            cell.fill = PatternFill("solid", fgColor=PHASE_COLOR[top])
    g.row_dimensions[r].height = 30
lr = 4 + len(ROWS) + 1
g.cell(row=lr, column=2, value="Marcos: 16/10 Business Case v2.0 | 06/11 formulário validado | 04/12 discovery consolidado | 11/12 Diagnóstico | 05/02/2027 apresentação final").font = Font(name="Calibri", size=10, italic=True, color="4A5560")
g.freeze_panes = "D4"
for sh in (g,):
    sh.page_setup.orientation = "landscape"
    sh.page_setup.fitToWidth = 1
    sh.sheet_properties.pageSetUpPr.fitToPage = True

# O que mudou
m = wb.create_sheet("O que mudou")
m.sheet_view.showGridLines = False
m.column_dimensions["A"].width = 30
m.column_dimensions["B"].width = 95
m["A1"] = "O que mudou em relação à versão anterior"
m["A1"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
CHANGES = [
    ("Repetições removidas", "\"Visita na Dreamer\" entrou no 3.1 (mesmas datas e objetivo do 3.1.1). \"Análise de mercado\" (4.2) já é coberta pelo benchmark (2.1.1) e pelo formulário (2.2). A SWOT (4.1) virou parte do Diagnóstico."),
    ("Ordem corrigida", "Primeiro pesquisar, depois decidir, depois escrever o que construir. O Diagnóstico (4.1) agora vem antes dos PRDs (5.1)."),
    ("Formulário reorganizado", "Rascunho feito por Rodrigo até 23/10. Revisão com a consultora Kátia Sato entre 26/10 e 06/11 (02/11 é feriado). Coleta de 09/11 a 27/11 e consolidação até 04/12."),
    ("Piloto separado", "O piloto saiu do Diagnóstico e virou fase própria (8.0). Um piloto de 50 fazendas não cabe em uma semana de dezembro, e muitos confinamentos pequenos e médios ficam vazios no fim do ano."),
    ("Um número só para o piloto", "A versão anterior tinha 10/20/50 e 50/100/500. Sugestão: começar com 5 fazendas e ampliar por etapas, com tamanho final definido no Roadmap (6.1)."),
    ("Ponto de atenção: Dreamer", "A entrega do TN Manager Web está prevista para o fim de novembro, antes do resultado da pesquisa (04/12). Sugestão: até lá, focar em correções e no que já existe, deixando funcionalidades novas para depois do Diagnóstico."),
]
for r, (a, b) in enumerate(CHANGES, 3):
    ca, cb = m.cell(row=r, column=1, value=a), m.cell(row=r, column=2, value=b)
    ca.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ca.fill = PatternFill("solid", fgColor=MAG if r % 2 else NAVY)
    ca.alignment = WRAP
    cb.font = Font(name="Calibri", size=11)
    cb.alignment = WRAP
    for c in (ca, cb):
        c.border = BORDER
    m.row_dimensions[r].height = 48

wb.save(OUT)
print("ok", OUT)
