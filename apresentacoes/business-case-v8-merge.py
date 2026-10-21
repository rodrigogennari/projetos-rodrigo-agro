import copy, re, sys, io
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
S=sys.argv[1]
dst=Presentation(S+'/v7/out.pptx'); src=Presentation(S+'/build/venc.pptx')
orig=list(dst.slides)
lay=[l for l in dst.slide_layouts if l.name=='DEFAULT'][0]
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
P='{http://schemas.openxmlformats.org/presentationml/2006/main}'
ss=src.slides[0]; d=dst.slides.add_slide(lay)
for ph in list(d.placeholders): ph._element.getparent().remove(ph._element)
for el in ss.shapes._spTree.iterchildren():
    if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
    e=copy.deepcopy(el)
    for b in e.iter():
        rid=b.get(R+'embed')
        if rid:
            _,nr=d.part.get_or_add_image_part(io.BytesIO(ss.part.related_part(rid).blob)); b.set(R+'embed',nr)
    d.shapes._spTree.append(e)
bg=ss._element.cSld.find(P+'bg')
if bg is not None: d._element.cSld.insert(0,copy.deepcopy(bg))
ns=d.notes_slide; tf=ns.notes_text_frame
if tf is not None: tf.text=ss.notes_slide.notes_text_frame.text
else:
    t=ns.shapes._spTree
    for el in list(t.iterchildren())[2:]: t.remove(el)
    for el in list(ss.notes_slide.shapes._spTree.iterchildren())[2:]: t.append(copy.deepcopy(el))
def title(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip(): return sh.text_frame.text
    return ''
k=[i for i,x in enumerate(orig) if title(x).startswith('Concorrência')][0]
order=orig[:k+1]+[d]+orig[k+1:]
lst=dst.slides._sldIdLst
el_by={el.get('id'):el for el in lst}
oid=[str(s.slide_id) for s in order]
for el in list(lst): lst.remove(el)
for x in oid: lst.append(el_by[x])
H=914400
tb=d.shapes.add_textbox(Emu(int(5.3*H)),Emu(int(0.07*H)),Emu(int(4.5*H)),Emu(int(0.2*H)))
tf=tb.text_frame; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
p=tf.paragraphs[0]; p.alignment=PP_ALIGN.RIGHT
r=p.add_run(); r.text='CAPÍTULO 4  ·  CONTRA QUEM'; r.font.size=Pt(8); r.font.bold=True; r.font.name='Calibri'; r.font.color.rgb=RGBColor.from_string('C8106E'); r._r.get_or_add_rPr().set('spc','150')
def set_para(p,text):
    runs=p.runs; runs[0].text=text
    for x in runs[1:]: x._r.getparent().remove(x._r)
def paras(s):
    for sh in s.shapes:
        if sh.has_text_frame:
            yield from sh.text_frame.paragraphs
        if sh.has_table:
            for c in sh.table.iter_cells(): yield from c.text_frame.paragraphs
E=[("Cargill, DSM, Ponta, Tecnobeef e outros","Concorrentes e onde vamos vencer: suporte e dados")]
done=set()
def add_after(p, text):
    np_=copy.deepcopy(p._p); p._p.addnext(np_)
    from pptx.text.text import _Paragraph
    q=_Paragraph(np_, p._parent); set_para(q,text); return q
for s in order[:28]:
    for p in list(paras(s)):
        t=p.text
        for a,b in E:
            if t.startswith(a): set_para(p,b); done.add(a)
        if t.startswith("Sem custo de licença para o cliente"):
            q=add_after(p,"Suporte excelente, com analista online"); add_after(q,"Dashboards criativos para tomada de decisão"); done.add('forcas')
        if t.startswith("Reduzir licenças nas fazendas sem automação"):
            add_after(p,"Concorrentes com suporte fraco e implantação cara"); done.add('oport')
print('feito:',done)
for i,s in enumerate(order,1):
    for sh in s.shapes:
        if sh.has_text_frame and re.fullmatch(r'\d{1,2}',sh.text_frame.text.strip()) and sh.left is not None and sh.left>8.5*H and sh.top>4.9*H and sh.text_frame.paragraphs[0].runs:
            set_para(sh.text_frame.paragraphs[0],str(i))
dst.save(S+'/v8/out.pptx'); print('ok',len(order))
