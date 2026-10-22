import copy, re, sys, io
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
S=sys.argv[1]
dst=Presentation(S+'/v8/out.pptx'); src=Presentation(S+'/build/v9.pptx')
orig=list(dst.slides)
lay=[l for l in dst.slide_layouts if l.name=='DEFAULT'][0]
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
P='{http://schemas.openxmlformats.org/presentationml/2006/main}'
new=[]
for ss in src.slides:
    d=dst.slides.add_slide(lay)
    for ph in list(d.placeholders): ph._element.getparent().remove(ph._element)
    for el in ss.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
        e=copy.deepcopy(el)
        for b in e.iter():
            rid=b.get(R+'embed')
            if rid:
                _,nr=d.part.get_or_add_image_part(io.BytesIO(ss.part.related_part(rid).blob)); b.set(R+'embed',nr)
        d.shapes._spTree.append(e)
    bg=ss._element.cSld.find(P+'bg')
    if bg is not None: d._element.cSld.insert(0,copy.deepcopy(bg))
    if ss.has_notes_slide:
        ns=d.notes_slide; tf=ns.notes_text_frame
        if tf is not None: tf.text=ss.notes_slide.notes_text_frame.text
        else:
            t=ns.shapes._spTree
            for el in list(t.iterchildren())[2:]: t.remove(el)
            for el in list(ss.notes_slide.shapes._spTree.iterchildren())[2:]: t.append(copy.deepcopy(el))
    new.append(d)
O=lambda i: orig[i-1]; N=lambda i: new[i-1]
order=[O(1),O(2),N(1),O(6),N(2),O(8),O(5),O(7)]+[N(i) for i in (3,4,5,6,7,8,9,10,11,12,13,14,15,16,17,18,19,20)]+[O(28),O(29),O(10)]+[O(i) for i in range(30,44)]
removed=[O(i) for i in (3,4,9,11,12,13,14,15,16,17,18,19,20,21,22,23,24,25,26,27)]
lst=dst.slides._sldIdLst
el_by={el.get('id'):el for el in lst}
oid=[str(s.slide_id) for s in order]; rid=[str(s.slide_id) for s in removed]
assert len(set(oid))==len(oid) and not set(oid)&set(rid) and len(oid)+len(rid)+0==len(orig)+len(new)-0, (len(oid),len(rid),len(orig),len(new))
for x in rid:
    dst.part.drop_rel(el_by[x].get(R+'id')); lst.remove(el_by[x])
for el in list(lst): lst.remove(el)
for x in oid: lst.append(el_by[x])
def set_para(p,text):
    runs=p.runs; runs[0].text=text
    for x in runs[1:]: x._r.getparent().remove(x._r)
def paras(s):
    for sh in s.shapes:
        if sh.has_text_frame: yield from sh.text_frame.paragraphs
        if sh.has_table:
            for c in sh.table.iter_cells(): yield from c.text_frame.paragraphs
H=914400
# breadcrumbs nos slides mantidos
for s in [O(5),O(6),O(7),O(8),O(10)]:
    for sh in list(s.shapes):
        if sh.has_text_frame and sh.text_frame.text.startswith('CAPÍTULO'): sh._element.getparent().remove(sh._element)
for s in [O(6),O(8),O(5),O(7)]:
    tb=s.shapes.add_textbox(Emu(int(5.3*H)),Emu(int(0.07*H)),Emu(int(4.5*H)),Emu(int(0.2*H)))
    tf=tb.text_frame; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.RIGHT
    r=p.add_run(); r.text='CAPÍTULO 1  ·  CONTEXTO E PROPOSTA'; r.font.size=Pt(8); r.font.bold=True; r.font.name='Calibri'; r.font.color.rgb=RGBColor.from_string('C8106E'); r._r.get_or_add_rPr().set('spc','150')
done=[]
for p in list(paras(O(6))):
    t=p.text
    if t.startswith('O mercado está evoluindo'):
        set_para(p,'O mercado e a concorrência estão acelerando no digital. O confinamento representa 19% da nossa Unidade de Negócios Beef.'); done.append('r1')
    elif t.startswith('Este projeto inovador'):
        set_para(p,'Proposta: um TN Manager Web simples para novos clientes pequenos e médios, com suporte, dashboards e equipe técnica como diferencial, para fidelizar e abrir novos clientes.'); done.append('r2')
    elif t.startswith('Em um cenário em constante'):
        p._p.getparent().remove(p._p); done.append('r3')
for p in paras(O(7)):
    t=p.text
    if t.startswith('Uso recorrente'): set_para(p,'Primeiro: fazendas usando toda semana e satisfeitas. Depois: Sales Profit'); done.append('m2')
    elif t.startswith('Sales Profit e market share'): set_para(p,'Sales Profit e participação de mercado'); done.append('m1')
for p in paras(O(29)):
    if p.text.startswith('Para usar se a diretoria perguntar:'):
        set_para(p,p.text.replace('perguntar: ','perguntar: Raio-X do confinamento, ')); done.append('anexo')
print('feito',done)
for i,s in enumerate(order,1):
    for sh in s.shapes:
        if sh.has_text_frame and re.fullmatch(r'\d{1,2}',sh.text_frame.text.strip()) and sh.left is not None and sh.left>8.5*H and sh.top>4.9*H and sh.text_frame.paragraphs[0].runs:
            set_para(sh.text_frame.paragraphs[0],str(i))
dst.save(S+'/v9/out.pptx'); print('ok',len(order))
