import copy, re, sys, io
from pptx import Presentation
from pptx.util import Emu
S=sys.argv[1]
dst=Presentation(S+'/v2in/in.pptx'); src=Presentation(S+'/build/board.pptx')
orig=list(dst.slides)  # presentation order
lay=[l for l in dst.slide_layouts if l.name=='DEFAULT'][0]
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
new=[]
for ss in src.slides:
    d=dst.slides.add_slide(lay)
    for ph in list(d.placeholders): ph._element.getparent().remove(ph._element)
    sp=d.shapes._spTree
    for el in ss.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
        e=copy.deepcopy(el)
        for b in e.iter():
            rid=b.get(R+'embed')
            if rid:
                blob=ss.part.related_part(rid).blob
                _,nr=d.part.get_or_add_image_part(io.BytesIO(blob))
                b.set(R+'embed',nr)
        sp.append(e)
    bg=ss._element.cSld.find('{http://schemas.openxmlformats.org/presentationml/2006/main}bg')
    if bg is not None: d._element.cSld.insert(0,copy.deepcopy(bg))
    if ss.has_notes_slide:
        ns=d.notes_slide; tf=ns.notes_text_frame
        if tf is not None: tf.text=ss.notes_slide.notes_text_frame.text
        else:
            t=ns.shapes._spTree
            for el in list(t.iterchildren())[2:]: t.remove(el)
            for el in list(ss.notes_slide.shapes._spTree.iterchildren())[2:]: t.append(copy.deepcopy(el))
    new.append(d)
O=lambda i: orig[i-1]; N=lambda i: new[i-1]
order=[O(1),N(1),O(3),O(4),N(14),N(15),N(13),N(2),N(3),N(4),N(5),N(6),N(7),N(8),N(9),N(10),N(11),O(31),N(16),N(12),O(5),O(6),O(7),O(9),O(10),O(12),O(13),O(15),O(16),O(25),O(28),O(30),O(2),O(8),O(11),O(18),O(19),O(27)]
removed=[O(i) for i in (14,17,20,21,22,23,24,26,29)]
lst=dst.slides._sldIdLst
ids={s.slide_id:el for el in lst for s in [None] if False}
el_by={}
for el in list(lst):
    el_by[el.get('id')]=el
oid=[str(s.slide_id) for s in order]; rid=[str(s.slide_id) for s in removed]
for x in rid:
    el=el_by[x]; dst.part.drop_rel(el.get(R+'id')); lst.remove(el)
for el in list(lst): lst.remove(el)
for x in oid: lst.append(el_by[x])

def set_para(p, text):
    runs=p.runs
    runs[0].text=text
    for r in runs[1:]: r._r.getparent().remove(r._r)
def find_para(slide, start):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                if p.text.startswith(start): return p
        if sh.has_table:
            for c in sh.table.iter_cells():
                for p in c.text_frame.paragraphs:
                    if p.text.startswith(start): return p
    raise Exception(start)
# slide 4: diferencial
s4=O(4); ps=[p for sh in s4.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs]
k=[i for i,p in enumerate(ps) if p.text.startswith('O QUE NOS DIFERENCIA')][0]
set_para(ps[k+1],'Não é a ferramenta: software de gestão virou item de prateleira. O diferencial é a entrega técnica: a TN acompanha toda a engorda, analisa os números toda semana, compara com o benchmark e transforma dados em decisão para o confinador. Sem custo de licença.')
s25=O(25)
set_para(find_para(s25,'80% das fazendas'),'80% das fazendas ativas após 90 dias; nota de recomendação média de 8 ou mais; 20 fazendas atendidas pelo Rodrigo com auditoria mensal. Gatilho para contratar analistas (1 para cada 20 fazendas)')
set_para(find_para(s25,'Pelo menos 5 fazendas'),'Pelo menos 5 fazendas pedindo integração com balança ou vagão; custo por fazenda abaixo da licença paga hoje (média de R$ 18,2 mil por ano)')
set_para(find_para(O(28),'Piloto limitado'),'Rodrigo atende até 20 fazendas nos 6 primeiros meses, com auditoria mensal; depois, 1 analista dedicado para cada 20 fazendas (G2).')
# numeração
H=Emu(914400)
for i,s in enumerate(order,1):
    for sh in s.shapes:
        if sh.has_text_frame and re.fullmatch(r'\d{1,2}',sh.text_frame.text.strip()) and sh.left is not None and sh.left>8.5*H and sh.top>4.9*H:
            if sh.text_frame.paragraphs[0].runs: set_para(sh.text_frame.paragraphs[0],str(i))
O(2)._element.set('show','0')
dst.save(S+'/v2in/out.pptx'); print('ok',len(order))
