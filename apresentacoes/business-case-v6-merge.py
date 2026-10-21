import copy, re, sys, io
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
S=sys.argv[1]
dst=Presentation(S+'/v6/in.pptx'); src=Presentation(S+'/build/story.pptx')
orig=list(dst.slides)
lay=[l for l in dst.slide_layouts if l.name=='DEFAULT'][0]
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
P='{http://schemas.openxmlformats.org/presentationml/2006/main}'
new=[]
for ss in src.slides:
    d=dst.slides.add_slide(lay)
    for ph in list(d.placeholders): ph._element.getparent().remove(ph._element)
    sp=d.shapes._spTree
    for el in ss.shapes._spTree.iterchildren():
        if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
        e=copy.deepcopy(el)
        for b in e.iter():
            rid=b.get(R+'embed')
            if rid:
                _,nr=d.part.get_or_add_image_part(io.BytesIO(ss.part.related_part(rid).blob)); b.set(R+'embed',nr)
        sp.append(e)
    bg=ss._element.cSld.find(P+'bg')
    if bg is not None: d._element.cSld.insert(0,copy.deepcopy(bg))
    if ss.has_notes_slide:
        ns=d.notes_slide; tf=ns.notes_text_frame
        if tf is not None: tf.text=ss.notes_slide.notes_text_frame.text
        else:
            t=ns.shapes._spTree
            for el in list(t.iterchildren())[2:]: t.remove(el)
            for el in list(ss.notes_slide.shapes._spTree.iterchildren())[2:]: t.append(copy.deepcopy(el))
    new.append(d)
O=lambda i: orig[i-1]; N=lambda i: new[i-1]
order=[O(1),N(1),N(2),O(3),O(5),O(4),O(6),O(7),O(8),N(3),O(9),O(10),O(11),O(12),O(13),N(5),O(15),O(16),N(6),N(7),N(4),N(8),N(9),N(10),O(22)]+[O(i) for i in range(23,39)]
removed=[O(i) for i in (2,14,17,18,19,20,21)]
lst=dst.slides._sldIdLst
el_by={el.get('id'):el for el in lst}
oid=[str(s.slide_id) for s in order]; rid=[str(s.slide_id) for s in removed]
for x in rid:
    dst.part.drop_rel(el_by[x].get(R+'id')); lst.remove(el_by[x])
for el in list(lst): lst.remove(el)
for x in oid: lst.append(el_by[x])
CAP={1:'DE ONDE VIEMOS',2:'PARA QUEM',3:'O QUE ENTREGAMOS',4:'CONTRA QUEM',5:'COMO CRESCEMOS',6:'ATÉ ONDE VAMOS'}
chap=[(N(2),1),(O(3),1),(O(5),1),(O(4),2),(O(6),2),(O(7),2),(O(8),2),(N(3),2),(O(9),2),(O(10),3),(O(11),3),(O(12),3),(O(13),3),(N(5),4),(O(15),5),(O(16),5),(N(6),5),(N(7),5),(N(4),5),(N(8),6),(N(9),6),(N(10),6)]
H=914400
for s,c in chap:
    tb=s.shapes.add_textbox(Emu(int(5.3*H)),Emu(int(0.07*H)),Emu(int(4.5*H)),Emu(int(0.2*H)))
    tf=tb.text_frame; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.RIGHT
    r=p.add_run(); r.text=f'CAPÍTULO {c}  ·  {CAP[c]}'
    r.font.size=Pt(8); r.font.bold=True; r.font.name='Calibri'
    r.font.color.rgb=RGBColor.from_string('C5D300' if s is N(10) else 'C8106E')
    r._r.get_or_add_rPr().set('spc','150')
PONTE={
 3:"Ponte: 'Em resumo, é isso que propomos. E o que muda em relação à primeira versão?'",
 5:"Ponte: 'A pergunta mudou: deixou de ser como defender a base e passou a ser para quem geramos valor. Quem é esse cliente?'",
 4:"Ponte: 'E quantos confinamentos como esse existem no Brasil?'",
 6:"Ponte: 'São muitas fazendas, com pouco volume cada. Em cabeças, o mercado se concentra nos grandes.'",
 7:"Ponte: 'E onde estão esses confinadores?'",
 8:"Ponte: 'Além de onde estão, importa quem está decidindo nessas fazendas.'",
 9:"Ponte: 'Os dados confirmam o público. Agora, o que entregamos para ele?'",
 10:"Ponte: 'Todos ganham, mas o valor não está na ferramenta.'",
 11:"Ponte: 'Como isso acontece na rotina da fazenda?'",
 12:"Ponte: 'E na prática: papel ou celular? Precisa de internet?'",
 13:"Ponte: 'E o que a concorrência oferece?'",
 15:"Ponte: 'Com essas fazendas no piloto, quanto custa e como cobramos?'",
 16:"Ponte: 'Pondo tudo junto, onde somos fortes e onde estamos expostos?'",
}
PN=[(N(6),"Ponte: 'Como levamos isso ao mercado?'"),(N(7),"Ponte: 'E o que mantém a fazenda com a gente depois do piloto? A comunidade.'"),(N(8),"Ponte: 'Esses são os próximos 15 meses. E a visão até 2030?'"),(N(9),"Ponte: 'Para começar, precisamos de algumas decisões hoje.'"),(N(10),"Fechamento sugerido: 'Começamos pelo começo: rotina, dado, análise e resultado. Com o grupo e o evento anual, a fazenda fica pela comunidade, e a TN ganha dados e fidelidade.'")]
miss=[]
def prepend(s,text):
    tf=s.notes_slide.notes_text_frame
    if tf is None: miss.append(text); return
    old=tf.text; tf.text=text+('\n\n'+old if old.strip() else '')
for i,t in PONTE.items(): prepend(O(i),t)
for s,t in PN:
    tf=s.notes_slide.notes_text_frame
    if tf is None: miss.append(t)
    else: tf.text=tf.text+'\n\n'+t
def set_para(p,text):
    runs=p.runs; runs[0].text=text
    for r in runs[1:]: r._r.getparent().remove(r._r)
for i,s in enumerate(order,1):
    for sh in s.shapes:
        if sh.has_text_frame and re.fullmatch(r'\d{1,2}',sh.text_frame.text.strip()) and sh.left is not None and sh.left>8.5*H and sh.top>4.9*H and sh.text_frame.paragraphs[0].runs:
            set_para(sh.text_frame.paragraphs[0],str(i))
for sh in O(12).shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            if p.text.startswith('Analistas dedicados: 1 para cada'): set_para(p,'Analistas online: 1 para cada 20 fazendas')
dst.save(S+'/v6/out.pptx'); print('ok',len(order),'sem notas:',len(miss)); [print(' ',m[:60]) for m in miss]
