import copy, re, sys, io
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
S=sys.argv[1]
dst=Presentation(S+'/v6/out.pptx'); src=Presentation(S+'/build/quem.pptx')
orig=list(dst.slides)
lay=[l for l in dst.slide_layouts if l.name=='DEFAULT'][0]
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
P='{http://schemas.openxmlformats.org/presentationml/2006/main}'
ss=src.slides[0]; d=dst.slides.add_slide(lay)
for ph in list(d.placeholders): ph._element.getparent().remove(ph._element)
for el in ss.shapes._spTree.iterchildren():
    if el.tag.endswith('}nvGrpSpPr') or el.tag.endswith('}grpSpPr'): continue
    e=copy.deepcopy(el)
    for b in e.iter():
        rid=b.get(R+'embed')
        if rid:
            _,nr=d.part.get_or_add_image_part(io.BytesIO(ss.part.related_part(rid).blob)); b.set(R+'embed',nr)
    d.shapes._spTree.append(e)
bg=ss._element.cSld.find(P+'bg')
if bg is not None: d._element.cSld.insert(0,copy.deepcopy(bg))
ns=d.notes_slide; tf=ns.notes_text_frame
if tf is not None: tf.text=ss.notes_slide.notes_text_frame.text
else:
    t=ns.shapes._spTree
    for el in list(t.iterchildren())[2:]: t.remove(el)
    for el in list(ss.notes_slide.shapes._spTree.iterchildren())[2:]: t.append(copy.deepcopy(el))
def title(s):
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip(): return sh.text_frame.text
    return ''
cover=[s for s in orig if title(s).startswith('BUSINESS CASE')][0]
cover._element.attrib.pop('show',None)
O=lambda i: orig[i-1]
order=[cover,O(1),O(2),O(3),d]+[s for s in orig[3:] if s is not cover]
lst=dst.slides._sldIdLst
el_by={el.get('id'):el for el in lst}
oid=[str(s.slide_id) for s in order]
for el in list(lst): lst.remove(el)
for x in oid: lst.append(el_by[x])
H=914400
tb=d.shapes.add_textbox(Emu(int(5.3*H)),Emu(int(0.07*H)),Emu(int(4.5*H)),Emu(int(0.2*H)))
tf=tb.text_frame; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
p=tf.paragraphs[0]; p.alignment=PP_ALIGN.RIGHT
r=p.add_run(); r.text='CAPÍTULO 1  ·  DE ONDE VIEMOS'; r.font.size=Pt(8); r.font.bold=True; r.font.name='Calibri'; r.font.color.rgb=RGBColor.from_string('C8106E'); r._r.get_or_add_rPr().set('spc','150')
def set_para(p,text):
    runs=p.runs; runs[0].text=text
    for x in runs[1:]: x._r.getparent().remove(x._r)
def paras(s):
    for sh in s.shapes:
        if sh.has_text_frame:
            yield from sh.text_frame.paragraphs
        if sh.has_table:
            for c in sh.table.iter_cells(): yield from c.text_frame.paragraphs
E=[("Automação, se o cliente pedir","Automação: PecCode hoje, TN Manager no futuro"),
("Confinamentos pequenos e médios, (200 a 2.000)","Novos clientes pequenos e médios (200 a 2.000 cabeças), sem automação no momento"),
("Confinamentos pequenos e médios com capacidade estática de 200 a 2.000","Novos clientes: confinamentos pequenos e médios, com capacidade estática de 200 a 2.000 cabeças, sem automação no momento e com gestão moderada, que fazem anotações em papel, planilhas ou sistema simples."),
("Hoje: licenças TGC e PecCode. Amanhã:","Web para novos clientes pequenos e médios; PecCode (parceira) para quem quer automação"),
("Hoje pagamos licença de terceiros para 20 fazendas.","Pagamos 20 licenças hoje. O Web atende novos clientes sem automação; a PecCode segue com os maiores."),
("Tudo que for web passa a ser TN Manager Web","Novos clientes pequenos e médios, sem automação: TN Manager Web"),
("Cada fazenda que migra deixa de custar","Fazendas atuais sem automação podem migrar e deixar de custar de R$ 8 mil a R$ 40 mil por ano em licença"),
("Cortar até R$ 364 mil por ano em licenças","Reduzir licenças nas fazendas sem automação"),
("Carteira com licença paga que cabe na faixa","Novos clientes de 200 a 2.000 cabeças, sem automação, indicados pelos técnicos e RTVs"),
("Integração com balança, vagão e NOA; uso sem sinal","Uso sem sinal e integração com NOA"),
("Licenças de terceiros","Automação"),
("Migrar o que for web para o TN Manager","Maiores seguem com a PecCode (parceria)"),
("PecCode só para automação","PecCode segue; preparar a retomada do TN Manager Automação"),
("TGC só onde houver integração","TN Manager Automação retomado, em paralelo com a Web"),
("Licença só para automação","Web e Automação lado a lado"),
("Aprovar o foco em confinamentos de 200 a 2.000 cabeças","Foco em novos clientes de 200 a 2.000 cabeças, sem automação; maiores seguem na PecCode"),
]
done=set()
for s in order[:27]:
    for p in paras(s):
        t=p.text
        if t.startswith("PecCode como parceira em automação"):
            p.runs[0].text="PecCode, parceria consolidada: "
            p.runs[1].text="clientes maiores que querem automação continuam com a PecCode, que está rodando bem. O TN Manager Automação volta no futuro, em paralelo com a Web."
            done.add('peccode'); continue
        for a,b in E:
            if t.startswith(a) and p.runs:
                set_para(p,b); done.add(a)
print('faltou:',[a for a,b in E if a not in done], 'peccode' in done)
NOTE={"Custo-benefício e cobrança":"Atualização: o TN Manager Web é para novos clientes pequenos e médios, sem automação no momento. Clientes maiores que querem automação continuam com a PecCode, parceria consolidada e rodando bem. O TN Manager Automação será retomado no futuro, em paralelo com a versão Web. A economia com licenças depende de quantas das 20 fazendas licenciadas não usam automação: levantar a lista.",
"Nossa história":"Atualização: deixar claro que a versão Web é para novos clientes, pequenos e médios, sem automação no momento; os maiores com automação seguem com a PecCode; o TN Manager Automação volta no futuro, em paralelo. Ponte: 'Então, quem atende quem?'"}
for s in order[:27]:
    for k,v in NOTE.items():
        if title(s).startswith(k):
            tf=s.notes_slide.notes_text_frame
            if tf is not None: tf.text=tf.text+'\n\n'+v
for i,s in enumerate(order,1):
    for sh in s.shapes:
        if sh.has_text_frame and re.fullmatch(r'\d{1,2}',sh.text_frame.text.strip()) and sh.left is not None and sh.left>8.5*H and sh.top>4.9*H and sh.text_frame.paragraphs[0].runs:
            set_para(sh.text_frame.paragraphs[0],str(i))
dst.save(S+'/v7/out.pptx'); print('ok',len(order))
